"""
generate_fixtures.py

One-off script that creates small synthetic sample files (.txt, .md, .pdf,
.docx) under tests/fixtures/, so the ingestion test can prove parsing works
for all four supported formats without needing real candidate documents yet.

Run once:
    .\.venv\Scripts\python.exe tests\generate_fixtures.py
"""

from pathlib import Path

from docx import Document as DocxDocument
from fpdf import FPDF

FIXTURES_DIR = Path(__file__).parent / "fixtures"

SAMPLE_JD_TEXT = (
    "Junior AI Engineer\n"
    "Northwind Labs, Karachi\n"
    "0-2 years experience\n\n"
    "Must-haves: Python fundamentals, an LLM framework such as LangChain or "
    "LangGraph, RAG including chunking and retrieval quality, Git and REST APIs, "
    "one shipped project."
)

SAMPLE_RESUME_TEXT = (
    "Jane Candidate\n"
    "Computer Science student, FAST NUCES\n\n"
    "Experience:\n"
    "Built a RAG chatbot using LangChain, ChromaDB, and Groq.\n"
    "Built an AI lead generation agent using Playwright and Streamlit.\n\n"
    "GitHub: https://github.com/example-candidate"
)


def create_txt():
    (FIXTURES_DIR / "sample_jd.txt").write_text(SAMPLE_JD_TEXT, encoding="utf-8")


def create_md():
    content = "# Job Description\n\n" + SAMPLE_JD_TEXT
    (FIXTURES_DIR / "sample_jd.md").write_text(content, encoding="utf-8")


def create_pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in SAMPLE_RESUME_TEXT.split("\n"):
        text = line if line.strip() != "" else " "
        pdf.cell(0, 8, text, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(FIXTURES_DIR / "sample_resume.pdf"))


def create_docx():
    doc = DocxDocument()
    for line in SAMPLE_RESUME_TEXT.split("\n"):
        doc.add_paragraph(line)
    doc.save(str(FIXTURES_DIR / "sample_resume.docx"))


if __name__ == "__main__":
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    create_txt()
    create_md()
    create_pdf()
    create_docx()
    print(f"Fixtures created in: {FIXTURES_DIR}")
    for f in sorted(FIXTURES_DIR.iterdir()):
        print(" -", f.name)